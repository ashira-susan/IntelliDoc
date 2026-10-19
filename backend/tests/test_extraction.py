from pathlib import Path

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.algorithms.extraction import (
    ExtractionError,
    extract_text,
)


def test_extract_txt(tmp_path: Path):
    path = tmp_path / "sample.txt"
    path.write_text(
        "Hello from Mr. IntelliDoc!",
        encoding="utf-8",
    )

    assert extract_text(path) == "Hello from Mr. IntelliDoc!"


def test_extract_docx_paragraphs_and_tables(tmp_path: Path):
    path = tmp_path / "sample.docx"

    document = DocxDocument()
    document.add_paragraph("IntelliDoc test document")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Subject"
    table.cell(0, 1).text = "Machine Learning"

    document.save(path)

    result = extract_text(path)

    assert "IntelliDoc test document" in result
    assert "Subject" in result
    assert "Machine Learning" in result


def test_extract_pdf(tmp_path: Path):
    path = tmp_path / "sample.pdf"

    with pymupdf.open() as document:
        page = document.new_page()
        page.insert_text(
            (72, 72),
            "Hello from IntelliDoc PDF",
        )
        document.save(path)

    result = extract_text(path)

    assert "Hello from IntelliDoc PDF" in result


def test_image_only_pdf_returns_empty_text(tmp_path: Path):
    path = tmp_path / "image_only.pdf"

    with pymupdf.open() as document:
        page = document.new_page()

        # Create a real image and insert it into the PDF.
        pixmap = pymupdf.Pixmap(
            pymupdf.csRGB,
            pymupdf.IRect(0, 0, 100, 100),
            0,
        )
        pixmap.clear_with(220)
        page.insert_image(page.rect, pixmap=pixmap)

        document.save(path)

    # The PDF contains an image but no selectable text.
    assert extract_text(path).strip() == ""


def test_corrupt_pdf_raises_extraction_error(tmp_path: Path):
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"This is not a valid PDF")

    with pytest.raises(ExtractionError):
        extract_text(path)


def test_unsupported_extension_raises_error(tmp_path: Path):
    path = tmp_path / "sample.csv"
    path.write_text("a,b", encoding="utf-8")

    with pytest.raises(ExtractionError):
        extract_text(path)
