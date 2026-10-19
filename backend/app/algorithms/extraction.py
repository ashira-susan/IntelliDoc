from pathlib import Path

import pymupdf
from docx import Document as DocxDocument


class ExtractionError(Exception):
    """Raised when a supported document cannot be read."""


def extract_txt(path: Path) -> str:
    """Read a text file using common encoding fallbacks."""
    content = path.read_bytes()

    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise ExtractionError("Could not decode the text file")


def extract_docx(path: Path) -> str:
    """Extract paragraphs and table cells from a Word document."""
    try:
        document = DocxDocument(path)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    except Exception as exc:
        raise ExtractionError(
            f"Could not read DOCX document: {exc}"
        ) from exc


def extract_pdf(path: Path) -> str:
    """Extract selectable text from PDF pages."""
    try:
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ExtractionError(
                    "Password-protected PDFs are not supported"
                )

            parts = [
                page.get_text()
                for page in document
            ]

            return "\n".join(parts)

    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            f"Could not read PDF document: {exc}"
        ) from exc


def extract_text(path: Path) -> str:
    """Dispatch extraction according to the file extension."""
    extension = path.suffix.lower()

    if extension == ".txt":
        return extract_txt(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension == ".pdf":
        return extract_pdf(path)

    raise ExtractionError(
        f"Unsupported file extension: {extension}"
    )
